"""Build the summarisation dataset from PriMock57 under the stage folder (EVAL_SUMMARISATION).

  prep/transcripts/<cid>.txt    all 57, "Doctor: ..." / "Patient: ..." lines by start time
  prep/checklists/<cid>.json    the 20 with a Clinician_1 checklist, as scorable items
  prep/manifest.json            per consult: transcript, checklist
  prep/human_labels.json        the clinicians' correct/incorrect labels on the dataset's
                                candidate notes, the validity reference for the judge
  notes/human-c<k>-note<n>/     those candidate notes as text, judged like any model's

Checklist roll-up: keep a top-level finding (no * prefix) or any critical leaf ("c"); drop
minor leaves and irrelevant ("i") items. Criticality c is critical, nc is minor.

    python eval/summarisation/prep.py
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from common import config, primock  # noqa: E402

LEVEL = re.compile(r"^(\*+)")
CRITICALITY = {"c": "critical", "nc": "minor"}
ARROW = "↳"  # sub-point marker in the note columns
NOTE_HEADER = re.compile(r"\s*note\s*(\d+)", re.I)


def checklist_dir():
    return config.path("primock57") / "consultation_checklists"


def consult_of(path):
    # "Day 1 Consultation 1 - Checklist.xlsx" -> day1_consultation01
    m = re.search(r"day\s*(\d+)\s*consultation\s*(\d+)", path.name, re.I)
    return f"day{int(m.group(1))}_consultation{int(m.group(2)):02d}" if m else None


def workbooks(clinician_dir):
    for path in sorted(clinician_dir.glob("*.xlsx")):
        if not path.name.startswith("~$") and consult_of(path):
            yield consult_of(path), path


def rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    out = list(wb.active.iter_rows(min_row=1, values_only=True))
    wb.close()
    return out


def checklist_items(sheet):
    items, seen = [], set()
    for row in sheet[1:]:
        if not row or row[0] is None:
            continue
        text = str(row[0]).strip()
        crit = str(row[1]).strip().lower() if len(row) > 1 and row[1] is not None else ""
        if crit not in ("c", "nc", "i") or not text:
            continue  # section headers and blanks
        m = LEVEL.match(text)
        level = len(m.group(1)) if m else 0
        clean = LEVEL.sub("", text).strip()
        if (level > 0 and crit != "c") or crit == "i":
            continue
        if clean.lower() in seen:
            continue
        seen.add(clean.lower())
        items.append({"id": f"c{len(items) + 1}", "text": clean, "criticality": CRITICALITY[crit]})
    return items


def note_columns(sheet):
    for row in sheet[:3]:
        cols = {int(NOTE_HEADER.match(str(c)).group(1)): j
                for j, c in enumerate(row) if c is not None and NOTE_HEADER.match(str(c))}
        if cols:
            return cols
    return {}


def note_text(sheet, col):
    lines = []
    for row in sheet:
        if col >= len(row) or row[col] is None:
            continue
        t = str(row[col]).strip()
        if not t or NOTE_HEADER.match(t):
            continue
        lines.append(t.replace(ARROW, "").strip())
    # Continuation fragments (lowercase start, or the line before ends mid-clause) join up
    out = []
    for t in lines:
        if out and (t[0].islower() or out[-1].rstrip().endswith((",", "and", "-"))):
            out[-1] = out[-1].rstrip() + " " + t
        else:
            out.append(t)
    return "\n".join(out).strip()


def note_labels(sheet, col):
    # The column after each note holds c (correct) or i (incorrect) per segment
    segments, correct, incorrect = [], 0, 0
    for row in sheet:
        if col >= len(row) or row[col] is None:
            continue
        t = str(row[col]).replace(ARROW, "").strip()
        if not t or NOTE_HEADER.match(t):
            continue
        mark = str(row[col + 1]).strip().lower() if col + 1 < len(row) and row[col + 1] is not None else ""
        if mark not in ("c", "i"):
            continue
        correct += mark == "c"
        incorrect += mark == "i"
        segments.append([t, mark == "c"])
    return segments, correct, incorrect


def main():
    data = config.path("summarisation")
    prep = data / "prep"
    (prep / "transcripts").mkdir(parents=True, exist_ok=True)
    (prep / "checklists").mkdir(parents=True, exist_ok=True)

    primary = checklist_dir() / "Clinician_1"
    sheets = {cid: rows(path) for cid, path in workbooks(primary)}
    manifest = []
    for cid in primock.consults():
        turns = [("Doctor" if s == "doctor" else "Patient", a, x) for a, _, s, x in primock.turns(cid)]
        turns.sort(key=lambda t: t[1])
        transcript = "\n".join(f"{who}: {text}" for who, _, text in turns)
        (prep / "transcripts" / f"{cid}.txt").write_text(transcript, encoding="utf-8")
        checklist = None
        if cid in sheets:
            checklist = f"checklists/{cid}.json"
            (prep / checklist).write_text(json.dumps(checklist_items(sheets[cid]), ensure_ascii=False, indent=2),
                                          encoding="utf-8")
        manifest.append({"consult_id": cid, "transcript": f"transcripts/{cid}.txt",
                         "checklist": checklist, "has_checklist": checklist is not None})
    (prep / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    labels = {}
    for cid, sheet in sheets.items():
        rec = {}
        for n, col in note_columns(sheet).items():
            segments, correct, incorrect = note_labels(sheet, col)
            if correct + incorrect:
                rec[f"note{n}"] = {"correct": correct, "incorrect": incorrect,
                                   "human_faith": round(correct / (correct + incorrect), 3),
                                   "n_seg": correct + incorrect, "segments": segments}
        if rec:
            labels[cid] = rec
    (prep / "human_labels.json").write_text(json.dumps(labels, ensure_ascii=False, indent=2), encoding="utf-8")

    notes = 0
    for clinician in sorted(checklist_dir().glob("Clinician_*")):
        who = clinician.name.lower().replace("clinician_", "c")
        for cid, path in workbooks(clinician):
            sheet = sheets[cid] if clinician == primary else rows(path)
            for n, col in note_columns(sheet).items():
                text = note_text(sheet, col)
                if len(text) < 40:  # column absent or empty
                    continue
                out = data / "notes" / f"human-{who}-note{n}"
                out.mkdir(parents=True, exist_ok=True)
                (out / f"{cid}.md").write_text(text, encoding="utf-8")
                notes += 1

    print(f"consults {len(manifest)}, with checklist {len(sheets)}, "
          f"labelled candidate notes {sum(len(v) for v in labels.values())}, human notes {notes}")


if __name__ == "__main__":
    main()
